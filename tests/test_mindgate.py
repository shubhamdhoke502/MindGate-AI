import gc
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch
from urllib.error import URLError

import pymupdf
from docx import Document
from fastapi.testclient import TestClient

from modules.app import MindGateEngine
from modules.chunk_cache import ChunkCacheManager
from modules.chat_history import ChatHistoryManager
from modules.parser import parse_pdf_production_grade
from modules.vector_db import VectorStoreManager
from ui import app_server


class ParserTests(unittest.TestCase):
    def test_pdf_chunks_have_page_and_extracted_line_coordinates(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "lines.pdf"
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((72, 72), "\n".join(
                    f"Evidence line {number}" for number in range(1, 11)
                ))
                pdf.save(path)

            chunks = parse_pdf_production_grade(str(path))

        self.assertEqual([chunk["metadata"]["line_start"] for chunk in chunks], [1, 7])
        self.assertEqual([chunk["metadata"]["line_end"] for chunk in chunks], [8, 10])
        self.assertTrue(all(chunk["metadata"]["page_number"] == 1 for chunk in chunks))
        self.assertIn("Evidence line 10", chunks[-1]["text"])
        self.assertEqual(chunks[0]["metadata"]["total_pages"], 1)
        self.assertGreater(chunks[0]["metadata"]["file_size_bytes"], 0)
        self.assertEqual(chunks[0]["metadata"]["language_detected"], "en")
        self.assertGreater(chunks[0]["metadata"]["character_count"], 0)
        self.assertTrue(chunks[0]["metadata"]["upload_timestamp"])

    def test_docx_ingestion_retains_unicode_paragraphs(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "notes.docx"
            document = Document()
            document.add_paragraph("मराठी माहिती")
            document.add_paragraph("Hindi जानकारी")
            document.save(path)

            chunks = parse_pdf_production_grade(str(path))

        self.assertEqual(len(chunks), 1)
        self.assertIn("मराठी", chunks[0]["text"])
        self.assertIn("जानकारी", chunks[0]["text"])
        self.assertEqual(chunks[0]["metadata"]["line_start"], 1)
        self.assertEqual(chunks[0]["metadata"]["line_end"], 2)
        self.assertEqual(chunks[0]["metadata"]["language_detected"], "hi")

    def test_marathi_script_is_detected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "marathi.txt"
            path.write_text("मळा आणि मराठी माहिती", encoding="utf-8")

            chunks = parse_pdf_production_grade(str(path))

        self.assertEqual(chunks[0]["metadata"]["language_detected"], "mr")


class RetrievalTests(unittest.TestCase):
    def test_bm25_is_local_and_never_escapes_requested_document(self):
        with tempfile.TemporaryDirectory() as directory:
            store = VectorStoreManager(storage_dir=directory)
            store.upsert_chunks([
                {
                    "chunk_id": "alpha",
                    "text": "The contract sets a royalty rate of five percent.",
                    "metadata": {
                        "file_name": "contract.pdf",
                        "page_number": 2,
                        "line_start": 3,
                        "line_end": 5,
                        "bbox": [1, 2, 3, 4],
                    },
                },
                {
                    "chunk_id": "beta",
                    "text": "The revenue report lists quarterly sales.",
                    "metadata": {"file_name": "report.pdf", "page_number": 1},
                },
            ])

            result = store.query_similar_chunks("royalty rate", file_name_filter="contract.pdf")
            missing = store.query_similar_chunks("quarterly sales", file_name_filter="missing.pdf")
            only_stop_words = store.query_similar_chunks("what is the")
            docs = store.list_documents()

            del store
            gc.collect()

        self.assertEqual([item["chunk_id"] for item in result], ["alpha"])
        self.assertEqual(result[0]["metadata"]["bbox"], [1, 2, 3, 4])
        self.assertEqual(missing, [])
        self.assertEqual(only_stop_words, [])
        self.assertEqual(docs, ["contract.pdf", "report.pdf"])

    def test_workspace_indexes_are_independent(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first = VectorStoreManager(storage_dir=str(root / "first"))
            second = VectorStoreManager(storage_dir=str(root / "second"))
            first.upsert_chunks([{
                "chunk_id": "same-name",
                "text": "private phrase in workspace one",
                "metadata": {"file_name": "notes.txt"},
            }])
            second.upsert_chunks([{
                "chunk_id": "same-name",
                "text": "different private phrase in workspace two",
                "metadata": {"file_name": "notes.txt"},
            }])

            first_result = first.query_similar_chunks("one", file_name_filter="notes.txt")
            second_result = second.query_similar_chunks("one", file_name_filter="notes.txt")
            del first, second
            gc.collect()

        self.assertEqual(len(first_result), 1)
        self.assertEqual(second_result, [])

    def test_reindex_replaces_stale_chunks_for_changed_document(self):
        with tempfile.TemporaryDirectory() as directory:
            store = VectorStoreManager(storage_dir=directory)
            store.replace_document("notes.txt", [{
                "chunk_id": "old",
                "text": "retired phrase",
                "metadata": {"file_name": "notes.txt"},
            }])
            store.replace_document("notes.txt", [{
                "chunk_id": "new",
                "text": "current phrase",
                "metadata": {"file_name": "notes.txt"},
            }])

            old = store.query_similar_chunks("retired", file_name_filter="notes.txt")
            current = store.query_similar_chunks("current", file_name_filter="notes.txt")
            store.close()

        self.assertEqual(old, [])
        self.assertEqual([item["chunk_id"] for item in current], ["new"])

    def test_hindi_and_marathi_terms_are_searchable(self):
        with tempfile.TemporaryDirectory() as directory:
            store = VectorStoreManager(storage_dir=directory)
            store.upsert_chunks([
                {
                    "chunk_id": "hindi",
                    "text": "ग्राहक को चौदह दिन में वापसी मिलेगी",
                    "metadata": {"file_name": "hindi.txt"},
                },
                {
                    "chunk_id": "marathi",
                    "text": "ग्राहकाला चौदा दिवसांत परतावा मिळेल",
                    "metadata": {"file_name": "marathi.txt"},
                },
            ])
            hindi = store.query_similar_chunks("चौदह दिन", file_name_filter="hindi.txt")
            marathi = store.query_similar_chunks("चौदा दिवसांत", file_name_filter="marathi.txt")
            store.close()

        self.assertEqual([item["chunk_id"] for item in hindi], ["hindi"])
        self.assertEqual([item["chunk_id"] for item in marathi], ["marathi"])

    def test_answer_archiving_supports_workspace_recall(self):
        with tempfile.TemporaryDirectory() as directory:
            engine = MindGateEngine(workspace_dir=directory)
            engine._archive_answer(
                "Refund requests are allowed within 14 days.",
                [{"file_name": "policy.pdf", "page": 2, "lines": "5-7"}],
            )
            recall = engine.container_memory.handle_at_trigger("refund")
            engine.vector_store.close()

        self.assertEqual(recall["matching_documents"], [])
        self.assertEqual(len(recall["matching_highlights"]), 1)
        self.assertEqual(recall["matching_highlights"][0]["page"], 2)
        self.assertEqual(recall["matching_highlights"][0]["line_range"], "5-7")

    def test_selected_document_is_never_replaced_with_global_search(self):
        class FakeVectorStore:
            def __init__(self):
                self.calls = []

            def query_similar_chunks(self, query, top_k, file_name_filter):
                self.calls.append(file_name_filter)
                return []

            def get_all_chunks_for_file(self, file_name, limit):
                return []

        engine = MindGateEngine.__new__(MindGateEngine)
        engine.vector_store = FakeVectorStore()
        engine.latest_uploaded_file = "other.pdf"

        result = engine.retrieve_relevant_chunks("specific question", "selected.pdf")

        self.assertEqual(result, [])
        self.assertEqual(engine.vector_store.calls, ["selected.pdf"])

    def test_cache_key_includes_source_name(self):
        with tempfile.TemporaryDirectory() as directory:
            first_path = Path(directory) / "first.txt"
            second_path = Path(directory) / "second.txt"
            first_path.write_text("identical", encoding="utf-8")
            second_path.write_text("identical", encoding="utf-8")
            cache = ChunkCacheManager(cache_dir=str(Path(directory) / "cache"))

            self.assertNotEqual(
                cache._get_file_hash(str(first_path)),
                cache._get_file_hash(str(second_path)),
            )

    def test_chat_history_serializes_concurrent_writes(self):
        with tempfile.TemporaryDirectory() as directory:
            history = ChatHistoryManager(workspace_dir=directory)
            with ThreadPoolExecutor(max_workers=8) as executor:
                list(executor.map(
                    lambda index: history.save_message("user", f"message {index}"),
                    range(40),
                ))
            messages = history.get_all_messages()

        self.assertEqual(len(messages), 40)

    def test_remote_ollama_endpoint_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "loopback"):
            MindGateEngine(ollama_host="https://example.com")

    def test_ollama_failure_is_not_returned_as_a_successful_answer(self):
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = workspace / "policy.txt"
            source.write_text("Refunds are available within fourteen days.", encoding="utf-8")
            engine = MindGateEngine(workspace_dir=str(workspace))
            engine.ingest_document(str(source))

            with patch("modules.app.urllib.request.urlopen", side_effect=URLError("offline")):
                with self.assertRaisesRegex(RuntimeError, "Ollama"):
                    engine.ask("How long are refunds available?", active_doc="policy.txt")
            engine.vector_store.close()


class ApiBoundaryTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.original_workspaces = app_server.WORKSPACES_DIR
        self.original_engines = app_server._engines
        app_server.WORKSPACES_DIR = Path(self.directory.name) / "workspaces"
        app_server._engines = {}
        self.client = TestClient(app_server.app, base_url="http://127.0.0.1")

    def tearDown(self):
        self.client.close()
        for engine in app_server._engines.values():
            engine.vector_store.close()
        app_server._engines = self.original_engines
        app_server.WORKSPACES_DIR = self.original_workspaces
        gc.collect()
        self.directory.cleanup()

    def test_workspace_traversal_and_unsupported_upload_are_rejected(self):
        traversal = self.client.post(
            "/workspace/init", data={"workspace_id": "..%5coutside"}
        )
        unsupported = self.client.post(
            "/upload",
            files={"file": ("payload.exe", b"not a document")},
        )

        self.assertEqual(traversal.status_code, 400)
        self.assertEqual(unsupported.status_code, 415)
        self.assertFalse((Path(self.directory.name).parent / "outside").exists())

    def test_document_count_and_pdf_page_limits_are_enforced(self):
        uploads = Path(self.directory.name) / "workspaces" / "limited" / "uploads"
        uploads.mkdir(parents=True)
        for index in range(app_server.MAX_WORKSPACE_DOCUMENTS):
            (uploads / f"existing-{index}.txt").write_text("existing", encoding="utf-8")

        count_limit = self.client.post(
            "/upload",
            files={"file": ("eleventh.txt", b"content")},
            data={"workspace_id": "limited"},
        )
        (uploads / "existing-9.txt").unlink()

        with pymupdf.open() as document:
            for _ in range(app_server.MAX_WORKSPACE_PDF_PAGES + 1):
                document.new_page()
            pdf_bytes = document.tobytes()
        page_limit = self.client.post(
            "/upload",
            files={"file": ("too-many-pages.pdf", pdf_bytes)},
            data={"workspace_id": "limited"},
        )

        self.assertEqual(count_limit.status_code, 413)
        self.assertEqual(page_limit.status_code, 413)

    def test_upload_size_limit_is_enforced(self):
        response = self.client.post(
            "/upload",
            files={
                "file": (
                    "too-large.txt",
                    b"x" * (app_server.MAX_UPLOAD_BYTES + 1),
                )
            },
            data={"workspace_id": "size_limit"},
        )

        self.assertEqual(response.status_code, 413)
        self.assertFalse(
            (Path(self.directory.name) / "workspaces" / "size_limit" / "uploads" / "too-large.txt").exists()
        )

    def test_cross_origin_mutations_are_rejected(self):
        response = self.client.post(
            "/workspace/init",
            data={"workspace_id": "isolated_test"},
            headers={"origin": "https://attacker.example"},
        )

        self.assertEqual(response.status_code, 403)

    def test_blank_chat_query_is_rejected(self):
        response = self.client.post("/chat", json={"query": "   "})

        self.assertEqual(response.status_code, 422)

    def test_document_upload_and_restore_list_work_without_ollama(self):
        response = self.client.post(
            "/upload",
            files={"file": ("notes.txt", b"Private workspace fact, number 381.")},
            data={"workspace_id": "isolated_test"},
        )
        previous_engine = app_server._engines.pop("isolated_test")
        previous_engine.vector_store.close()
        listing = self.client.get("/documents", params={"workspace_id": "isolated_test"})
        restored_chunks = app_server._engines["isolated_test"].vector_store.query_similar_chunks(
            "381", file_name_filter="notes.txt"
        )
        answer = self.client.post(
            "/chat",
            json={
                "query": "not present in the document",
                "workspace_id": "isolated_test",
                "active_pdf": "notes.txt",
            },
        )
        streamed = self.client.post(
            "/chat/stream",
            json={
                "query": "absent topic",
                "workspace_id": "isolated_test",
                "active_pdf": "notes.txt",
            },
        )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(listing.status_code, 200)
        self.assertEqual(listing.json()["documents"][0]["file_name"], "notes.txt")
        self.assertEqual(len(restored_chunks), 1)
        self.assertEqual(answer.status_code, 200, answer.text)
        self.assertEqual(
            answer.json()["result"]["answer"],
            "Is document me is question ka context nahi mila.",
        )
        history = self.client.get("/history", params={"workspace_id": "isolated_test"})
        self.assertEqual(history.status_code, 200)
        history_messages = history.json()["history"]
        self.assertEqual(len(history_messages), 4)
        self.assertTrue(all(message["file_name"] == "notes.txt" for message in history_messages))
        self.assertEqual(streamed.status_code, 200)
        self.assertIn("Is document me is question ka context nahi mila.", streamed.text)

    def test_pdf_upload_indexes_and_serves_inline_for_citations(self):
        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_text((72, 72), "The agreement allows a 21-day review period.")
            pdf_bytes = document.tobytes()

        upload = self.client.post(
            "/upload",
            files={"file": ("agreement.pdf", pdf_bytes)},
            data={"workspace_id": "pdf_test"},
        )
        document_response = self.client.get(
            "/documents/agreement.pdf",
            params={"workspace_id": "pdf_test"},
        )
        engine = app_server._engines["pdf_test"]
        retrieved = engine.vector_store.query_similar_chunks(
            "21-day review period", file_name_filter="agreement.pdf"
        )

        self.assertEqual(upload.status_code, 200, upload.text)
        self.assertEqual(upload.json()["ingest_result"]["total_chunks_indexed"], 1)
        self.assertEqual(document_response.status_code, 200)
        self.assertEqual(document_response.headers["content-type"], "application/pdf")
        self.assertEqual(document_response.headers["content-disposition"], "inline")
        self.assertTrue(document_response.content.startswith(b"%PDF-"))
        self.assertEqual(retrieved[0]["metadata"]["page_number"], 1)


if __name__ == "__main__":
    unittest.main()
